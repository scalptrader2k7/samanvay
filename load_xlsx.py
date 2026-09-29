"""Read text-cell XLSX uploads with explicit sheet/header selection.

This conservative adapter rejects formulas and non-text cells with coordinates.
It never guesses lost identifier digits, evaluates formulas or changes input.
"""
import hashlib
import io
from pathlib import Path

from openpyxl import load_workbook

from load_corpus import _map_rows, validate_mapping


def load_mapped_xlsx(path, mapping, *, sheet, header_row, import_root=None):
    validate_mapping(mapping)
    if not isinstance(sheet, str) or not sheet.strip():
        raise ValueError("An explicit worksheet name is required")
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        raise ValueError("An explicit positive header_row is required")
    path = Path(path).resolve()
    root = Path(import_root).resolve() if import_root is not None else path.parent
    try:
        source_file = path.relative_to(root).as_posix()
    except ValueError as exc:
        raise ValueError("Input file must be inside import_root") from exc
    snapshot = path.read_bytes()
    # data_only=False is essential: cached values must never hide formulas.
    workbook = load_workbook(io.BytesIO(snapshot), read_only=True, data_only=False)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet}")
        ws = workbook[sheet]
        if header_row > ws.max_row:
            raise ValueError("Header row is outside the worksheet")

        def text_cell(cell):
            if cell.data_type == 'f':
                raise ValueError(f"{sheet}!{cell.coordinate}: formula cell is unsupported; provide source text")
            if cell.data_type == 'e' or (cell.value is not None and not isinstance(cell.value, str)):
                raise ValueError(f"{sheet}!{cell.coordinate}: non-text cell is unsupported (type={cell.data_type}, format={cell.number_format}); numeric identifiers may have lost digits; provide explicit text")
            return cell.value if cell.value is not None else ''

        iterator = ws.iter_rows(min_row=header_row)
        headers = [text_cell(cell) for cell in next(iterator)]
        if not headers or any(not h.strip() for h in headers):
            raise ValueError("Missing or empty worksheet headers")
        if len(set(headers)) != len(headers):
            raise ValueError("Duplicate worksheet headers")
        needed = set(mapping['columns'].values()) | set(mapping.get('identifiers', [])) | set(mapping.get('specifications', []))
        missing = needed - set(headers)
        if missing:
            raise ValueError(f"Missing mapped columns: {', '.join(sorted(missing))}")
        rows = [dict(zip(headers, [text_cell(cell) for cell in cells])) for cells in iterator]
    finally:
        workbook.close()
    constant = mapping.get('organization_constant')
    records = _map_rows(rows, source_file, hashlib.sha256(snapshot).hexdigest(),
                        mapping['columns'], mapping.get('identifiers', []), mapping['mapping_id'],
                        (constant['value'], constant['evidence']) if constant else None,
                        mapping.get('specifications', []), mapping.get('legacy_code_evidence'))
    for record in records:
        record['source_sheet'] = sheet
        record['source_row'] += header_row
    return records
