"""Load a CSV/XLSX and emit a readiness/sample pack into a NEW directory."""
import argparse
from collections import Counter
import csv
import hashlib
import io
import json
from pathlib import Path

from input_profiles import PROFILES, get_profile
from load_corpus import SCALAR_FIELDS, load_mapped_csv, read_mapping


def _json(path, value):
    with path.open('x', encoding='utf-8', newline='\n') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')


def readiness(records, headers, mapping, source):
    """Completeness and validation counts, never model accuracy or identity counts."""
    fields = {}
    for target in SCALAR_FIELDS:
        column = mapping['columns'].get(target)
        constant = target == 'source_organization' and mapping.get('organization_constant')
        populated = sum(r[target] is not None for r in records)
        fields[target] = {
            'origin': 'column' if column else ('constant' if constant else 'unmapped'),
            'source_column': column,
            'column_present': column in headers if column else False,
            'populated_rows': populated, 'missing_rows': len(records) - populated,
        }
    for target, key in [('source_identifiers', 'identifiers'), ('specifications_original', 'specifications')]:
        columns = mapping.get(key, [])
        populated = sum(any(v.strip() for v in r[target].values()) for r in records)
        fields[target] = {'mapped_columns': list(columns),
                          'columns_present': [c for c in columns if c in headers],
                          'populated_rows': populated, 'missing_rows': len(records) - populated}
    issue_counts = Counter((i['field'], i['reason'], i['severity']) for r in records for i in r['issues'])
    eligible = sum(r['eligible_for_processing'] for r in records)
    description_counts = Counter(r['description_original'] for r in records if r['description_original'] is not None)
    return {
        'schema_version': '0.1', 'source': source, 'mapping_id': mapping['mapping_id'],
        'input_rows': len(records), 'output_records': len(records), 'dropped_rows': 0,
        'eligible_rows': eligible, 'ineligible_rows': len(records) - eligible,
        'rows_with_issues': sum(bool(r['issues']) for r in records),
        'repeated_description_rows_beyond_first': sum(n - 1 for n in description_counts.values()),
        'source_columns': headers, 'fields': fields,
        'issues': [dict(field=f, reason=r, severity=s, count=n)
                   for (f, r, s), n in sorted(issue_counts.items())],
        'readiness': {
            'legacy_codes': 'missing' if not fields['legacy_material_code']['populated_rows'] else 'supplied_not_verified',
            'specifications': 'missing' if not fields['specifications_original']['populated_rows'] else 'supplied_not_verified',
            'transaction_history': 'not_assessed_no_transaction_adapter',
            'transaction_modules_enabled': False,
            'human_validation': 'not_assessed_by_ingestion',
        },
        'limitations': [
            'Counts describe input rows, not unique materials or verified matches.',
            'Overlapping sources and candidate subsets must not be summed as distinct materials.',
            'No PR/PO/stock/issue history is mapped. Current public inputs do not provide it.',
            'Planned quantities and estimated values do not establish consumption, prices paid or savings.',
            'Human review and sponsor material-master validation remain separate from ingestion.',
        ],
    }


def ingest(path, mapping, output_dir, *, import_root=None, sheet=None, header_row=None, sample_size=5):
    """Validate fully, then exclusively create a pack; never overwrite any output.

    A write failure can leave an incomplete directory. manifest.json is written
    last and marks completion; rerun into a different new directory if incomplete.
    """
    if isinstance(sample_size, bool) or not isinstance(sample_size, int) or sample_size < 1:
        raise ValueError('sample_size must be a positive integer')
    path, output = Path(path), Path(output_dir)
    if output.exists():
        raise FileExistsError(f'Output directory already exists: {output}')
    before = path.read_bytes()
    digest = hashlib.sha256(before).hexdigest()
    if path.suffix.lower() == '.csv':
        if sheet is not None or header_row is not None:
            raise ValueError('sheet/header_row apply only to XLSX')
        records = load_mapped_csv(path, mapping, import_root=import_root)
        headers = next(csv.reader(io.StringIO(before.decode('utf-8-sig'), newline='')))
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        from load_xlsx import load_mapped_xlsx
        records = load_mapped_xlsx(path, mapping, import_root=import_root, sheet=sheet, header_row=header_row)
        workbook = load_workbook(io.BytesIO(before), read_only=True, data_only=False)
        try:
            headers = [c.value for c in next(workbook[sheet].iter_rows(min_row=header_row, max_row=header_row))]
        finally:
            workbook.close()
    else:
        raise ValueError('Only .csv and .xlsx inputs are supported')
    if path.read_bytes() != before or any(r['source_sha256'] != digest for r in records):
        raise ValueError('Input changed during import; no output written')
    root = Path(import_root).resolve() if import_root is not None else path.resolve().parent
    source = {'source_file': path.resolve().relative_to(root).as_posix(), 'source_sha256': digest,
              'source_sheet': sheet, 'header_row': header_row}
    report = readiness(records, headers, mapping, source)
    output.mkdir(parents=True, exist_ok=False)
    _json(output / 'mapping.json', mapping)
    _json(output / 'readiness.json', report)
    with (output / 'records.jsonl').open('x', encoding='utf-8', newline='\n') as stream:
        for record in records:
            stream.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + '\n')
    sample = records[:sample_size]
    with (output / 'sample.csv').open('x', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(r['raw_fields'] for r in sample)
    _json(output / 'sample_provenance.json', [
        {'sample_row': index, **{k: r[k] for k in ('source_file', 'source_sha256', 'source_sheet', 'source_row', 'source_record_id', 'source_identifiers')}}
        for index, r in enumerate(sample, 1)])
    lines = ['# Input readiness', '', f"Source: `{source['source_file']}`",
             f"Mapping: `{mapping['mapping_id']}`", '',
             f"Loaded {len(records)} rows; retained {len(records)} records; dropped 0.",
             f"Eligible: {report['eligible_rows']}; ineligible: {report['ineligible_rows']}.", '',
             '| Field | Origin | Populated rows | Missing rows |', '|---|---|---:|---:|']
    for field, values in report['fields'].items():
        origin = values.get('origin', 'columns' if values.get('mapped_columns') else 'unmapped')
        lines.append(f"| {field} | {origin} | {values['populated_rows']} | {values['missing_rows']} |")
    lines.extend(['', 'No transaction adapter is present; transaction modules are disabled.', '',
                  *report['limitations'], '',
                  'The sample contains original cells from the first input rows, not invented examples.',
                  'Use sample_provenance.json to trace each sample back to the original snapshot.',
                  'records.jsonl includes each retained row and its validation issues.'])
    with (output / 'readiness.md').open('x', encoding='utf-8', newline='\n') as stream:
        stream.write('\n'.join(lines) + '\n')
    artifacts = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(output.iterdir())}
    _json(output / 'manifest.json', {'status': 'complete', 'source': source,
                                    'output_records': len(records), 'sample_rows': len(sample),
                                    'artifact_sha256': artifacts})
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input')
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--profile', choices=sorted(PROFILES))
    group.add_argument('--mapping', help='Explicit mapping JSON')
    parser.add_argument('--output-dir', required=True)
    parser.add_argument('--import-root')
    parser.add_argument('--sheet')
    parser.add_argument('--header-row', type=int)
    parser.add_argument('--sample-size', type=int, default=5)
    args = parser.parse_args()
    try:
        mapping = get_profile(args.profile) if args.profile else read_mapping(args.mapping)
        report = ingest(args.input, mapping, args.output_dir, import_root=args.import_root,
                        sheet=args.sheet, header_row=args.header_row, sample_size=args.sample_size)
    except (ValueError, OSError) as exc:
        parser.exit(1, f'Import failed: {exc}\n')
    print(f"Loaded {report['input_rows']} rows; pack: {args.output_dir}")


if __name__ == '__main__':
    main()
