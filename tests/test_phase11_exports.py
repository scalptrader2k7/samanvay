from io import BytesIO
from pathlib import Path
import csv
import tempfile
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api_service import Settings, create_app
from prototype_registry import load_phase3_records
from service_storage import LocalWorkspaceStorage


ROOT = Path(__file__).resolve().parents[1]


class Phase11ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        records = {r["source_record_id"]: r for r in load_phase3_records(
            ROOT / "output/phase3-v0.1/material_attributes.jsonl")}
        cls.records = [records[x] for x in ("M000753", "M000925", "M000830")]

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.storage = LocalWorkspaceStorage(Path(self.temp.name) / "storage")
        self.storage.create_workspace("export-demo", "export-owner", public_read=True)
        self.storage.create_workspace("export-private", "other-owner", public_read=False)
        with self.storage.open_workspace("export-demo", "export-owner", write=True) as store:
            store.import_records(self.records)
            store.decide("approve", "identity", ["M000753", "M000925"],
                         "Supported matching pipe evidence", "export-owner")
        self.client = TestClient(create_app(settings=Settings(
            demo_token="export-secret", demo_subject="export-owner"), storage=self.storage))
        self.auth = {"Authorization": "Bearer export-secret"}
        self.base = "/api/v1/workspaces/export-demo/exports"

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def test_registry_and_mapping_exports_preserve_current_ids_and_paginate(self):
        registry_json = self.client.get(f"{self.base}/registry.json?limit=1").json()
        self.assertEqual(registry_json["total"], 1)
        row = registry_json["items"][0]
        self.assertTrue(row["proposed_code"].startswith("SMV-P-"))
        self.assertEqual(row["member_source_ids"], ["M000753", "M000925"])
        first = self.client.get(f"{self.base}/source-mappings.csv?limit=1")
        second = self.client.get(f"{self.base}/source-mappings.csv?limit=1")
        self.assertEqual(first.content, second.content)
        mapping = next(csv.DictReader(first.text.splitlines()))
        self.assertEqual(mapping["source_record_id"], "M000753")

    def test_excel_pack_is_reproducible_current_state_and_separates_history(self):
        first = self.client.get(f"{self.base}/quality-pack.xlsx")
        second = self.client.get(f"{self.base}/quality-pack.xlsx")
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.content, second.content)
        self.assertEqual(first.headers["x-export-scope"], "current-state-no-history")
        workbook = load_workbook(BytesIO(first.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames,
                         ["Manifest", "Dashboard", "Completeness", "Registry", "Current Mappings"])
        registry_rows = list(workbook["Registry"].iter_rows(values_only=True))
        self.assertIn("M000753", registry_rows[1][6])
        mapping_rows = list(workbook["Current Mappings"].iter_rows(values_only=True))
        source_index = mapping_rows[0].index("source_record_id")
        self.assertEqual({row[source_index] for row in mapping_rows[1:]},
                         {"M000753", "M000925", "M000830"})
        self.assertNotIn("Audit Events", workbook.sheetnames)
        history = self.client.get(f"{self.base}/audit-events.json").json()
        self.assertEqual(history["total"], 1)

    def test_export_workspace_isolation(self):
        response = self.client.get(
            "/api/v1/workspaces/export-private/exports/quality-pack.xlsx",
            headers=self.auth)
        self.assertEqual(response.status_code, 403)


if __name__ == "__main__":
    unittest.main()
